from __future__ import annotations

import json
import re
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(r"C:\Users\동아대도서관\Documents\카카오톡 받은 파일\2. 취업지원실_도서목록(2026.9.3)나눔 행사용 목록.xlsx")
DATA_DIR = ROOT / "public" / "assets" / "data"
OUTPUT_DIR = ROOT / "outputs" / "career-book-giveaway"


def clean(value) -> str:
    return str(value or "").strip()


def price(value) -> int:
    digits = re.sub(r"[^0-9]", "", clean(value))
    return int(digits) if digits else 0


def author(value) -> str:
    return re.sub(r"\s+", " ", clean(value)).strip(" ,")


def category(title: str, call_no: str) -> str:
    text = f"{title} {call_no}".lower()
    rules = [
        ("면접/스피치", ["면접", "interview", "스피치", "speech", "말하기"]),
        ("자기소개서/이력서", ["자소서", "자기소개", "이력서", "resume", "입사지원서"]),
        ("NCS/인적성", ["ncs", "인적성", "직무적성", "적성", "psat", "gsat", "cat", "cjat", "sat"]),
        ("공무원/자격시험", ["공무원", "공기업", "자격", "기사", "한국사", "토익", "toeic", "컴활"]),
        ("창업/경영경제", ["창업", "경영", "마케팅", "회계", "재무", "경제", "기업"]),
        ("직무/실무역량", ["직무", "실무", "엑셀", "excel", "python", "파이썬", "ai", "데이터", "디자인", "포트폴리오"]),
        ("진로/취업전략", ["취업", "진로", "커리어", "채용", "직업", "합격"]),
    ]
    for label, words in rules:
        if any(word in text for word in words):
            return label
    kdc_match = re.search(r"(\d{3})", call_no)
    if kdc_match:
        code = int(kdc_match.group(1))
        if 300 <= code <= 399:
            if 320 <= code <= 329:
                return "창업/경영경제"
            if 330 <= code <= 339:
                return "진로/취업전략"
            return "사회과학/취업교양"
        if 000 <= code <= 99:
            return "컴퓨터/데이터"
        if 100 <= code <= 199:
            return "심리/자기계발"
        if 400 <= code <= 499:
            return "과학/기술"
        if 500 <= code <= 599:
            return "기술/실무"
        if 600 <= code <= 699:
            return "예술/디자인"
        if 700 <= code <= 799:
            return "어학/글쓰기"
        if 800 <= code <= 899:
            return "문학/교양"
        if 900 <= code <= 999:
            return "역사/인문"
    return "취업교양/기타"


def read_books() -> list[dict]:
    wb = load_workbook(SOURCE, data_only=True, read_only=True)
    ws = wb.active
    books = []
    for row in ws.iter_rows(min_row=3, values_only=True):
        if not row or not row[0]:
            continue
        source_no, reg_no, title, call_no, writer, location, shelf_date, acquired_date, raw_price, year, detail_url = row[:11]
        title = clean(title)
        reg_no = clean(reg_no)
        if not title or not reg_no:
            continue
        call_no = clean(call_no)
        books.append({
            "bookId": f"BOOK-{len(books) + 1:05d}",
            "sourceNo": int(source_no) if isinstance(source_no, (int, float)) else clean(source_no),
            "registrationNo": reg_no,
            "title": title,
            "author": author(writer),
            "callNo": call_no,
            "location": clean(location),
            "shelfDate": clean(shelf_date),
            "acquisitionDate": clean(acquired_date),
            "price": price(raw_price),
            "publicationYear": clean(year),
            "detailUrl": clean(detail_url),
            "isbn13": "",
            "category": category(title, call_no),
            "status": "신청가능",
            "availableQuantity": 1,
            "note": "",
        })
    return books


def write_json(books: list[dict]) -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    payload = {
        "generatedAt": datetime.now().isoformat(timespec="seconds"),
        "sourceFile": SOURCE.name,
        "count": len(books),
        "books": books,
    }
    (DATA_DIR / "career-books.json").write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")


def header_style(ws, row: int, cols: int, fill: str) -> None:
    for cell in ws[row][:cols]:
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def widths(ws, values: list[float]) -> None:
    for index, width in enumerate(values, 1):
        ws.column_dimensions[get_column_letter(index)].width = width


def row_style(ws, max_row: int, max_col: int) -> None:
    border = Border(bottom=Side(style="thin", color="DDE4F0"))
    for row in ws.iter_rows(min_row=2, max_row=max_row, max_col=max_col):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)


def write_template(books: list[dict]) -> Path:
    wb = Workbook()
    ws_books = wb.active
    ws_books.title = "도서목록"
    ws_requests = wb.create_sheet("신청현황")
    ws_users = wb.create_sheet("이용자")
    ws_settings = wb.create_sheet("설정")

    book_headers = ["도서ID", "등록번호", "서명", "저자", "청구기호", "소장위치", "가격", "출판년도", "도서상세URL", "ISBN13", "카테고리", "상태", "신청가능수량", "신청중수량", "확정수량", "비고", "원본번호"]
    request_headers = ["신청ID", "신청일시", "상태", "학생명", "학번", "학과", "연락처", "이메일", "신청경로", "도서ID", "등록번호", "서명", "저자", "ISBN13", "메모", "처리자", "처리일시"]
    user_headers = ["학번", "성명", "휴대폰번호", "개인정보동의", "최초로그인", "최근로그인", "로그인횟수"]

    ws_books.append(book_headers)
    for book in books:
        row_index = ws_books.max_row + 1
        ws_books.append([book["bookId"], book["registrationNo"], book["title"], book["author"], book["callNo"], book["location"], book["price"], book["publicationYear"], book["detailUrl"], book["isbn13"], book["category"], "", 1, "", "", book["note"], book["sourceNo"]])
        ws_books.cell(row_index, 12).value = f'=IF($M{row_index}<=($N{row_index}+$O{row_index}),"마감","신청가능")'
        ws_books.cell(row_index, 14).value = f'=COUNTIFS(신청현황!$J:$J,$A{row_index},신청현황!$C:$C,"신청접수")+COUNTIFS(신청현황!$J:$J,$A{row_index},신청현황!$C:$C,"처리중")'
        ws_books.cell(row_index, 15).value = f'=COUNTIFS(신청현황!$J:$J,$A{row_index},신청현황!$C:$C,"확정")'

    ws_requests.append(request_headers)
    ws_users.append(user_headers)
    ws_settings.append(["설정항목", "값", "비고"])
    for key, note in [
        ("SPREADSHEET_ID", "이 구글시트의 ID"),
        ("ALADIN_TTB_KEY", "알라딘 TTBKey는 Apps Script 프로젝트 속성에 저장 권장"),
        ("EVENT_STATUS", "OPEN"),
        ("WEB_APP_URL", "Apps Script 배포 후 웹 앱 URL 입력"),
    ]:
        ws_settings.append([key, "", note])

    for ws in (ws_books, ws_requests, ws_users, ws_settings):
        ws.freeze_panes = "A2"
        ws.sheet_view.showGridLines = False

    header_style(ws_books, 1, len(book_headers), "1559A8")
    header_style(ws_requests, 1, len(request_headers), "E23B78")
    header_style(ws_users, 1, len(user_headers), "17A765")
    header_style(ws_settings, 1, 3, "263238")
    widths(ws_books, [14, 14, 44, 22, 18, 18, 12, 10, 44, 16, 15, 12, 12, 12, 12, 24, 10])
    widths(ws_requests, [24, 20, 12, 14, 14, 18, 16, 24, 14, 14, 14, 44, 22, 16, 28, 14, 20])
    widths(ws_users, [16, 14, 18, 14, 20, 20, 12])
    widths(ws_settings, [22, 36, 58])
    row_style(ws_books, ws_books.max_row, len(book_headers))
    row_style(ws_requests, 5000, len(request_headers))
    row_style(ws_users, 3000, len(user_headers))
    row_style(ws_settings, ws_settings.max_row, 3)

    ws_books.auto_filter.ref = f"A1:Q{ws_books.max_row}"
    ws_requests.auto_filter.ref = "A1:Q5000"
    ws_users.auto_filter.ref = "A1:G3000"
    ws_requests.add_data_validation(DataValidation(type="list", formula1='"신청접수,처리중,확정,취소,중복,마감"', allow_blank=False))
    ws_requests.data_validations.dataValidation[-1].add("C2:C5000")
    ws_books.conditional_formatting.add(f"A2:Q{ws_books.max_row}", FormulaRule(formula=['$L2="마감"'], fill=PatternFill("solid", fgColor="FDE2E2")))
    ws_requests.conditional_formatting.add("A2:Q5000", FormulaRule(formula=['$C2="신청접수"'], fill=PatternFill("solid", fgColor="FFF4CC")))

    for row in range(2, ws_books.max_row + 1):
        ws_books.cell(row, 7).number_format = '#,##0'

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    output = OUTPUT_DIR / "dong-a-career-book-giveaway-google-sheets-template.xlsx"
    wb.save(output)
    return output


def main() -> None:
    books = read_books()
    write_json(books)
    output = write_template(books)
    print(json.dumps({"books": len(books), "template": str(output)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
